"""
YES24 베스트셀러 EDA 프레젠테이션 네오브루탈리즘(Neo-Brutalism) 스타일 변환 스크립트

주요 역할:
- 기존 노르딕 스타일 30페이지 PPTX(create_presentation.py 산출물)를 읽어 내용·차트·발표자 노트는 그대로 두고
  디자인만 네오브루탈리즘으로 재스타일링합니다.
- 네오브루탈리즘 규격(pptx-design-styles 02번):
  · 고채도 단색 배경(옐로 #F5F500), 순백 카드, 순검정 테두리(2.5~3pt)
  · 블러 없는 하드 오프셋 그림자(우하단), 둥근 모서리 제거
  · 액센트 컬러 블루 #0000FF / 레드 #FF3B30, 표지·간지는 블랙/블루 단색 배경
  · 간지와 표지에 레이아웃을 깨는 초대형 숫자/스티커 요소 배치
- 폰트는 사용자 지정(제목: G마켓 산스 Bold, 본문: 나눔고딕)을 유지하며,
  한글 글리프에 실제 적용되도록 동아시아(ea) 서체 속성까지 런 단위로 명시합니다.

작성자: Antigravity Senior Data Analyst
생성일: 2026년 10월
"""

import os
import copy

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL, MSO_COLOR_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

SRC = os.path.join("yes24", "docs", "yes24_day_bestseller_eda_presentation.pptx")
DST = os.path.join("yes24", "docs", "yes24_day_bestseller_eda_presentation_neobrutalism.pptx")

FONT_TITLE = "Gmarket Sans TTF Bold"
FONT_BODY = "NanumGothic"

# 네오브루탈리즘 팔레트
YELLOW = "F5F500"
LIME = "CCFF00"
BLACK = "000000"
WHITE = "FFFFFF"
BLUE = "0000FF"
RED = "FF3B30"
PINK = "FF2D55"
MAGENTA = "E6007A"
ROW_ALT = "FFF9C4"

# 원본(노르딕) 색상 → 역할 판별용
NAVY_BG = (24, 35, 44)
SLATE_BG = (28, 43, 54)
LIGHT_BLUE = (226, 236, 244)
BORDER_GREY = (226, 230, 233)

# 원본 텍스트/바 색상 → 네오브루탈리즘 색상 매핑 (콘텐츠 슬라이드)
CONTENT_COLOR_MAP = {
    (36, 42, 48): BLACK,      # 노르딕 차콜 본문
    (100, 116, 139): "222222",  # 슬레이트 보조 텍스트
    (46, 90, 136): BLUE,      # 포인트 블루
    (46, 92, 62): MAGENTA,    # 세이지 그린
    (148, 91, 45): RED,       # 테라코타
    (24, 35, 44): BLACK,      # 딥 네이비
}

SHADOW_OFFSET = Inches(0.09)


def rgb_tuple(color_format):
    """ColorFormat에서 RGB 튜플을 안전하게 읽습니다 (미지정 시 None)."""
    try:
        if color_format is not None and color_format.type == MSO_COLOR_TYPE.RGB:
            c = color_format.rgb
            return (c[0], c[1], c[2])
    except (AttributeError, TypeError):
        pass
    return None


def fill_rgb(shape):
    """도형의 단색 채우기 RGB 튜플을 반환합니다."""
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return rgb_tuple(shape.fill.fore_color)
    except (AttributeError, TypeError):
        pass
    return None


def set_fill(shape, hex_color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)


def set_line(shape, hex_color=None, width_pt=2.5):
    if hex_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(hex_color)
        shape.line.width = Pt(width_pt)


def square_corners(shape):
    """둥근 사각형을 직각 사각형으로 변경합니다 (네오브루탈리즘은 둥근 모서리 금지)."""
    geom = shape._element.spPr.find(qn("a:prstGeom"))
    if geom is not None and geom.get("prst") == "roundRect":
        geom.set("prst", "rect")
        av = geom.find(qn("a:avLst"))
        if av is not None:
            for gd in list(av):
                av.remove(gd)


def add_hard_shadow(slide, shape, color=BLACK, offset=SHADOW_OFFSET):
    """블러 없는 하드 오프셋 그림자 사각형을 대상 도형 바로 뒤(z-order)에 삽입합니다."""
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, shape.left + offset, shape.top + offset, shape.width, shape.height
    )
    set_fill(sh, color)
    set_line(sh, None)
    sh.rotation = getattr(shape, "rotation", 0) or 0
    el = sh._element
    el.getparent().remove(el)
    shape._element.addprevious(el)
    return sh


def apply_font_to_run(run, name, size, bold, hex_color):
    """런 단위로 서체(latin/ea/cs)·크기·굵기·색상을 명시합니다."""
    f = run.font
    if hex_color:
        f.color.rgb = RGBColor.from_string(hex_color)
    if size is not None:
        f.size = size
    if bold is not None:
        f.bold = bold
    if name:
        f.name = name
        rPr = run._r.get_or_add_rPr()
        latin = rPr.find(qn("a:latin"))
        anchor = latin
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
                el.getparent().remove(el)
                if anchor is not None:
                    anchor.addnext(el)
                else:
                    rPr.append(el)
            el.set("typeface", name)
            anchor = el


def restyle_text_frame(tf, color_fn, force_color=None):
    """텍스트 프레임의 모든 문단 색상을 재매핑하고 서체 속성을 런에 명시합니다."""
    for p in tf.paragraphs:
        pf = p.font
        orig = rgb_tuple(pf.color) if pf.color is not None else None
        if orig is None and p.runs:
            orig = rgb_tuple(p.runs[0].font.color)
        new_color = force_color or color_fn(orig)
        name = pf.name or (p.runs[0].font.name if p.runs else None)
        size = pf.size
        bold = pf.bold
        if new_color:
            pf.color.rgb = RGBColor.from_string(new_color)
        for r in p.runs:
            apply_font_to_run(
                r,
                name or r.font.name,
                size or r.font.size,
                bold if bold is not None else r.font.bold,
                new_color,
            )


def content_color(orig):
    if orig is None:
        return BLACK
    return CONTENT_COLOR_MAP.get(orig, BLACK)


def dark_slide_color(orig):
    """표지/간지 슬라이드의 텍스트: 흰색 유지, 밝은 블루 계열은 옐로로 치환."""
    if orig is None or orig == (255, 255, 255):
        return WHITE
    return YELLOW


def set_cell_border(cell, hex_color=BLACK, width_emu=19050):
    """표 셀 4면에 검정 테두리를 지정합니다 (tcPr 내 ln* 요소는 채우기보다 앞에 위치해야 함)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for tag in reversed(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.SubElement(tcPr, qn(tag))
        tcPr.remove(ln)
        ln.set("w", str(width_emu))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr"))
        clr.set("val", hex_color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(0, ln)


def restyle_table(graphic_frame):
    table = graphic_frame.table
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = RGBColor.from_string(BLACK)
                restyle_text_frame(cell.text_frame, content_color, force_color=YELLOW)
            else:
                cell.fill.fore_color.rgb = RGBColor.from_string(ROW_ALT if r_idx % 2 == 0 else WHITE)
                restyle_text_frame(cell.text_frame, content_color, force_color=BLACK)
            set_cell_border(cell)


def add_text_sticker(slide, left, top, width, height, text, fill, text_color, size_pt,
                     rotation=0, shadow_color=BLACK, line_color=BLACK):
    """네오브루탈리즘 스티커(굵은 테두리 + 하드 그림자 + 기울기) 텍스트 블록을 추가합니다."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    set_fill(box, fill)
    set_line(box, line_color, 3)
    box.rotation = rotation
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        apply_font_to_run(r, FONT_TITLE, Pt(size_pt if i == 0 else max(12, size_pt * 0.38)), True, text_color)
    if shadow_color:
        add_hard_shadow(slide, box, shadow_color)
    return box


def add_giant_number(slide, left, top, width, height, text, hex_color):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    apply_font_to_run(r, FONT_TITLE, Pt(150), True, hex_color)
    return tb


def restyle_dark_slide(slide, bg_shape, kind, section_no=None):
    """표지·클로징(kind='black') 및 간지(kind='blue') 슬라이드 재스타일링."""
    set_fill(bg_shape, BLACK if kind == "black" else BLUE)
    shadow_color = WHITE if kind == "black" else BLACK

    for shape in list(slide.shapes):
        if shape is bg_shape or shape._element is bg_shape._element:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            square_corners(shape)
            is_oval = shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL
            if is_oval:
                set_fill(shape, PINK)
                set_line(shape, None)
                continue
            if shape.height < Inches(0.1):  # 얇은 구분선
                set_fill(shape, YELLOW)
                set_line(shape, None)
                shape.height = Pt(5)
                continue
            if shape.height < Inches(0.5):  # 배지
                if kind == "black":
                    set_fill(shape, YELLOW)
                    txt = BLACK
                else:
                    set_fill(shape, BLACK)
                    txt = YELLOW
                set_line(shape, None)
                shape.rotation = -3
                if shape.has_text_frame:
                    restyle_text_frame(shape.text_frame, dark_slide_color, force_color=txt)
                add_hard_shadow(slide, shape, shadow_color, Inches(0.05))
                continue
            # 큰 카드/하이라이트 박스
            set_fill(shape, YELLOW if kind == "black" else WHITE)
            set_line(shape, BLACK, 3)
            if shape.has_text_frame:
                restyle_text_frame(shape.text_frame, dark_slide_color, force_color=BLACK)
            add_hard_shadow(slide, shape, shadow_color)
        elif shape.has_text_frame:
            restyle_text_frame(shape.text_frame, dark_slide_color)

    if kind == "blue" and section_no is not None:
        add_giant_number(slide, Inches(8.8), Inches(5.35), Inches(4.1), Inches(2.0),
                         f"{section_no:02d}", YELLOW)


def restyle_content_slide(slide, bg_shape):
    set_fill(bg_shape, YELLOW)

    for shape in list(slide.shapes):
        if shape._element is bg_shape._element:
            continue
        st = shape.shape_type
        if st == MSO_SHAPE_TYPE.PICTURE:
            shape.line.color.rgb = RGBColor.from_string(BLACK)
            shape.line.width = Pt(3)
            add_hard_shadow(slide, shape)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            restyle_table(shape)
            add_hard_shadow(slide, shape)
            continue
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            square_corners(shape)
            orig_fill = fill_rgb(shape)

            # 1) 헤더 하단 얇은 구분선 → 굵은 검정 라인
            if shape.width > Inches(10) and shape.height <= Pt(2):
                set_fill(shape, BLACK)
                set_line(shape, None)
                shape.height = Pt(5)
                continue
            # 2) 카드 상단/좌측 액센트 바 → 두껍고 선명하게
            if shape.height <= Pt(5) or shape.width <= Pt(5):
                accent = CONTENT_COLOR_MAP.get(orig_fill, BLUE) if orig_fill else BLUE
                if accent in (BLACK, "222222"):
                    accent = BLACK
                set_fill(shape, accent)
                set_line(shape, BLACK, 2)
                if shape.height <= Pt(5):
                    shape.height = Pt(9)
                else:
                    shape.width = Pt(12)
                continue
            # 3) 헤더 배지 → 검정 배지 + 옐로 텍스트, 살짝 기울임
            if shape.height < Inches(0.45):
                set_fill(shape, BLACK)
                set_line(shape, None)
                shape.rotation = -2
                if shape.has_text_frame:
                    restyle_text_frame(shape.text_frame, content_color, force_color=YELLOW)
                add_hard_shadow(slide, shape, BLUE, Inches(0.05))
                continue
            # 4) 하이라이트 박스(아이스 블루) → 라임
            if orig_fill == LIGHT_BLUE:
                set_fill(shape, LIME)
            else:
                set_fill(shape, WHITE)
            set_line(shape, BLACK, 2.5)
            if shape.has_text_frame:
                restyle_text_frame(shape.text_frame, content_color)
            add_hard_shadow(slide, shape)
            continue
        if shape.has_text_frame:
            restyle_text_frame(shape.text_frame, content_color)


def find_background(slide, prs):
    for shape in slide.shapes:
        if (shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.left == 0 and shape.top == 0
                and shape.width == prs.slide_width and shape.height == prs.slide_height):
            return shape
    return None


def main():
    if not os.path.exists(SRC):
        # 원본 노르딕 덱이 없으면 먼저 생성
        from create_presentation import build_presentation  # noqa: E402
        build_presentation()

    prs = Presentation(SRC)
    section_no = 0
    for idx, slide in enumerate(prs.slides):
        bg = find_background(slide, prs)
        if bg is None:
            continue
        bg_rgb = fill_rgb(bg)
        if bg_rgb == NAVY_BG:
            restyle_dark_slide(slide, bg, "black")
            if idx == 0:
                # 표지: 레이아웃을 깨는 기울어진 스티커
                add_text_sticker(slide, Inches(10.0), Inches(0.55), Inches(2.6), Inches(1.35),
                                 "285건\n전수 분석 · 10/10 PASS", YELLOW, BLACK, 30,
                                 rotation=6, shadow_color=WHITE)
        elif bg_rgb == SLATE_BG:
            section_no += 1
            restyle_dark_slide(slide, bg, "blue", section_no)
        else:
            restyle_content_slide(slide, bg)

    prs.save(DST)
    print(f"네오브루탈리즘 덱 저장 완료: {DST} (슬라이드 {len(prs.slides)}장)")


if __name__ == "__main__":
    main()
